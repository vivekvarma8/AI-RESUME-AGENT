import PyPDF2
import docx
import pdfplumber
from pathlib import Path

class ResumeParser:
    """Extracts text from PDF and DOCX resume files"""
    
    def __init__(self):
        pass
    
    def parse(self, file_path):
        """
        Main method to parse resume
        Returns: extracted text as string
        """
        file_path = Path(file_path)
        
        if not file_path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")
        
        extension = file_path.suffix.lower()
        
        if extension == '.pdf':
            return self._parse_pdf(file_path)
        elif extension == '.docx':
            return self._parse_docx(file_path)
        else:
            raise ValueError(f"Unsupported file type: {extension}")
    
    def _parse_pdf(self, file_path):
        """Extract text from PDF using pdfplumber (better formatting)"""
        text = ""
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    page_text = page.extract_text()
                    if page_text:
                        text += page_text + "\n"
            
            # If pdfplumber fails, try PyPDF2
            if not text.strip():
                text = self._parse_pdf_pypdf2(file_path)
                
        except Exception as e:
            # Fallback to PyPDF2
            print(f"pdfplumber failed, trying PyPDF2: {e}")
            text = self._parse_pdf_pypdf2(file_path)
        
        return text.strip()
    
    def _parse_pdf_pypdf2(self, file_path):
        """Fallback PDF parser using PyPDF2"""
        text = ""
        with open(file_path, 'rb') as file:
            pdf_reader = PyPDF2.PdfReader(file)
            for page in pdf_reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text += page_text + "\n"
        return text.strip()
    
    def _parse_docx(self, file_path):
        """Extract text from DOCX"""
        doc = docx.Document(file_path)
        text = ""
        
        for paragraph in doc.paragraphs:
            text += paragraph.text + "\n"
        
        # Also extract text from tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    text += cell.text + " "
                text += "\n"
        
        return text.strip()